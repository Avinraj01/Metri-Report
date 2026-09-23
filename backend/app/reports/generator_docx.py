import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.models import Report, TestSession, Instrument

def generate_docx_report(report: Report, output_path: str) -> str:
    """
    Generates an editable DOCX format OIML R-76 test report.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    session: TestSession = report.session
    instrument: Instrument = session.instrument

    # Title & Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_header.add_run("GOVERNMENT OF INDIA\nMINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION\nDEPARTMENT OF CONSUMER AFFAIRS — LEGAL METROLOGY DIVISION\n")
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(15, 23, 42)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p_title.add_run("TYPE EVALUATION TEST REPORT\n")
    r2.bold = True
    r2.font.size = Pt(15)
    r2.font.color.rgb = RGBColor(30, 58, 138)
    r3 = p_title.add_run("Non-Automatic Weighing Instrument as per OIML R 76-1:2006 (E)\n")
    r3.font.size = Pt(10)
    r3.italic = True

    # Summary Table
    table_meta = doc.add_table(rows=3, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.rows[0].cells[0].text = f"Report Number: {report.report_number}"
    table_meta.rows[0].cells[1].text = f"Date: {datetime.utcnow().strftime('%d-%b-%Y')}"
    table_meta.rows[1].cells[0].text = f"Session Number: {session.session_number}"
    overall_val = report.overall_result.value if hasattr(report.overall_result, 'value') else str(report.overall_result)
    overall_disp = "MANUAL REVIEW / INCOMPLETE" if overall_val == "MANUAL_REVIEW" else overall_val
    table_meta.rows[2].cells[0].text = f"Overall Result: {overall_disp}"
    table_meta.rows[2].cells[1].text = "Standard: OIML R 76-1:2006 (E)"
    doc.add_paragraph()

    # Section 1: Applicant & Manufacturer
    h1 = doc.add_heading("1. Applicant & Manufacturer Information", level=2)
    h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    t1 = doc.add_table(rows=4, cols=2)
    t1.rows[0].cells[0].text = "Applicant Name"
    t1.rows[0].cells[1].text = instrument.applicant_name
    t1.rows[1].cells[0].text = "Applicant Address"
    t1.rows[1].cells[1].text = instrument.applicant_address or "N/A"
    t1.rows[2].cells[0].text = "Manufacturer Name"
    t1.rows[2].cells[1].text = instrument.manufacturer
    t1.rows[3].cells[0].text = "Manufacturer Address"
    t1.rows[3].cells[1].text = instrument.manufacturer_address or "N/A"
    doc.add_paragraph()

    # Section 2: Technical Characteristics
    h2 = doc.add_heading("2. Technical & Metrological Characteristics", level=2)
    h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    t2 = doc.add_table(rows=6, cols=2)
    t2.rows[0].cells[0].text = "Model & Serial Number"
    t2.rows[0].cells[1].text = f"{instrument.model} / S/N: {instrument.serial_number}"
    t2.rows[1].cells[0].text = "Accuracy Class"
    t2.rows[1].cells[1].text = f"Class {instrument.accuracy_class.value}"
    t2.rows[2].cells[0].text = "Max & Min Capacity"
    t2.rows[2].cells[1].text = f"Max = {instrument.max_capacity} {instrument.unit}, Min = {instrument.min_capacity} {instrument.unit}"
    t2.rows[3].cells[0].text = "Scale Intervals (e, d, n)"
    t2.rows[3].cells[1].text = f"e = {instrument.e_value} {instrument.unit}, d = {instrument.d_value} {instrument.unit}, n = {instrument.n_intervals}"
    t2.rows[4].cells[0].text = "Load Receptor / Supports"
    t2.rows[4].cells[1].text = f"{instrument.load_receptor_type} ({instrument.num_support_points} support points)"
    t2.rows[5].cells[0].text = "Power & Temp Range"
    t2.rows[5].cells[1].text = f"{instrument.power_supply} | {instrument.temp_min}°C to {instrument.temp_max}°C"
    doc.add_paragraph()

    # Section 3: Test Applicability Matrix
    from sqlalchemy.orm import object_session
    from app.models import TestCatalogItem
    from app.rules.applicability import ApplicabilityEngine

    db = object_session(report)
    catalog_items = db.query(TestCatalogItem).filter(TestCatalogItem.is_active == True).all() if db else []
    applicable_matrix = ApplicabilityEngine.get_applicable_tests(instrument, catalog_items) if catalog_items else []
    executed_codes = {cr.test_code: cr for cr in session.compliance_results}

    h3_app = doc.add_heading("3. OIML R-76 Test Applicability & Coverage Matrix", level=2)
    h3_app.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    
    t_app = doc.add_table(rows=len(applicable_matrix) + 1, cols=6)
    app_headers = ["Code", "Clause", "Test Procedure", "Applicability", "Execution", "Status"]
    for i, h in enumerate(app_headers):
        cell = t_app.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True

    for row_idx, item in enumerate(applicable_matrix, start=1):
        is_app = item["is_applicable"]
        t_code = item["test_code"]
        is_exec = t_code in executed_codes
        
        exec_str = "EXECUTED" if is_exec else ("N/A" if not is_app else "NOT TESTED")
        if is_exec:
            status_str = executed_codes[t_code].status.value
        elif not is_app:
            status_str = "NOT APPLICABLE"
        else:
            status_str = "PENDING"

        t_app.rows[row_idx].cells[0].text = t_code
        t_app.rows[row_idx].cells[1].text = item["clause"]
        t_app.rows[row_idx].cells[2].text = item["test_name"]
        t_app.rows[row_idx].cells[3].text = "APPLICABLE" if is_app else "EXEMPT / N/A"
        t_app.rows[row_idx].cells[4].text = exec_str
        t_app.rows[row_idx].cells[5].text = status_str
    doc.add_paragraph()

    # Section 4: Executed Test Compliance Summary
    h4 = doc.add_heading("4. Summary of Executed Test Results & Metrological Compliance", level=2)
    h4.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    comp_results = session.compliance_results
    t4 = doc.add_table(rows=len(comp_results) + 1, cols=5)
    headers = ["Clause", "Test Description", "Observed Result", "Permissible Limit", "Result"]
    for i, h in enumerate(headers):
        cell = t4.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True

    for row_idx, cr in enumerate(comp_results, start=1):
        t4.rows[row_idx].cells[0].text = cr.clause
        t4.rows[row_idx].cells[1].text = cr.test_name
        t4.rows[row_idx].cells[2].text = cr.observed_summary
        t4.rows[row_idx].cells[3].text = cr.permissible_summary
        t4.rows[row_idx].cells[4].text = cr.status.value
    doc.add_paragraph()

    # Section 5: Signatures
    p_sig = doc.add_paragraph("\n5. Review & Approval Signatures\n")
    p_sig.runs[0].bold = True
    t_sig = doc.add_table(rows=1, cols=3)
    t_sig.rows[0].cells[0].text = f"Evaluated by:\n\n___________________\nTest Engineer\nDate: {datetime.utcnow().strftime('%d-%b-%Y')}"
    
    rev_text = "Reviewed by:\n\n___________________\nTechnical Reviewer\n"
    if report.status.value in ["APPROVED", "FINALIZED"]:
        rev_text += f"Status: {report.status.value}"
    elif report.status.value == "CHANGES_REQUESTED":
        rev_text += "Status: CHANGES REQUESTED"
    elif report.status.value in ["READY_FOR_REVIEW", "UNDER_REVIEW"]:
        rev_text += f"Status: {report.status.value.replace('_', ' ')} (Pending Review)"
    else:
        rev_text += "Status: DRAFT"
    t_sig.rows[0].cells[1].text = rev_text

    app_text = "Approved by:\n\n___________________\nDirector / Lab In-Charge\n"
    if report.status.value == "FINALIZED":
        final_date = report.finalized_at.strftime('%d-%b-%Y') if report.finalized_at else ''
        app_text += f"Status: FINALIZED / SEALED\nDate: {final_date}"
    elif report.status.value == "APPROVED":
        app_text += "Status: APPROVED (Pending Final Seal)"
    else:
        app_text += "Status: Pending Final Approval"
    t_sig.rows[0].cells[2].text = app_text

    # Disclaimer
    p_disc = doc.add_paragraph("\nDISCLAIMER: This prototype test report is generated by METRIREPORT for SIH 2026. It is not an official model approval certificate under the Legal Metrology Act.")
    p_disc.runs[0].font.size = Pt(8)
    p_disc.runs[0].font.italic = True

    doc.save(output_path)
    return output_path
